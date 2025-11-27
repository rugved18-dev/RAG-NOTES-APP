"""Document processing utilities for loading and embedding multiple file types."""
import os
from typing import List, Tuple
from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.schema import Document


def load_pdf(file_path: str) -> Tuple[List[Document], str]:
    """
    Load PDF and extract text with metadata about page numbers.
    
    Args:
        file_path: Path to PDF file
        
    Returns:
        Tuple of (list of Document objects, file name)
    """
    file_name = os.path.basename(file_path)
    documents = []
    
    try:
        with open(file_path, 'rb') as pdf_file:
            pdf_reader = PdfReader(pdf_file)
            num_pages = len(pdf_reader.pages)
            
            for page_num in range(num_pages):
                page = pdf_reader.pages[page_num]
                text = page.extract_text()
                
                # Create document with metadata including page number and file name
                doc = Document(
                    page_content=text,
                    metadata={
                        "source": file_name,
                        "page": page_num + 1,
                        "total_pages": num_pages,
                        "file_type": "PDF"
                    }
                )
                documents.append(doc)
                
        return documents, file_name
    except Exception as e:
        raise Exception(f"Error loading PDF {file_name}: {str(e)}")


def load_docx(file_path: str) -> Tuple[List[Document], str]:
    """
    Load Word document and extract text with paragraph metadata.
    
    Args:
        file_path: Path to DOCX file
        
    Returns:
        Tuple of (list of Document objects, file name)
    """
    file_name = os.path.basename(file_path)
    documents = []
    
    try:
        doc = DocxDocument(file_path)
        para_num = 0
        
        for para in doc.paragraphs:
            if para.text.strip():
                document = Document(
                    page_content=para.text,
                    metadata={
                        "source": file_name,
                        "paragraph": para_num + 1,
                        "file_type": "DOCX"
                    }
                )
                documents.append(document)
                para_num += 1
        
        # Also extract from tables
        for table_num, table in enumerate(doc.tables):
            for row_num, row in enumerate(table.rows):
                row_text = " | ".join(cell.text for cell in row.cells)
                if row_text.strip():
                    document = Document(
                        page_content=row_text,
                        metadata={
                            "source": file_name,
                            "table": table_num + 1,
                            "row": row_num + 1,
                            "file_type": "DOCX"
                        }
                    )
                    documents.append(document)
        
        return documents, file_name
    except Exception as e:
        raise Exception(f"Error loading DOCX {file_name}: {str(e)}")


def load_pptx(file_path: str) -> Tuple[List[Document], str]:
    """
    Load PowerPoint presentation and extract text with slide metadata.
    
    Args:
        file_path: Path to PPTX file
        
    Returns:
        Tuple of (list of Document objects, file name)
    """
    file_name = os.path.basename(file_path)
    documents = []
    
    try:
        presentation = Presentation(file_path)
        
        for slide_num, slide in enumerate(presentation.slides):
            slide_text = []
            
            # Extract text from shapes
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)
            
            # Combine all text from slide
            combined_text = "\n".join(slide_text)
            
            if combined_text.strip():
                document = Document(
                    page_content=combined_text,
                    metadata={
                        "source": file_name,
                        "slide": slide_num + 1,
                        "file_type": "PPTX"
                    }
                )
                documents.append(document)
        
        return documents, file_name
    except Exception as e:
        raise Exception(f"Error loading PPTX {file_name}: {str(e)}")


def chunk_documents(documents: List[Document], chunk_size: int = 1000, 
                   chunk_overlap: int = 200) -> List[Document]:
    """
    Split documents into chunks for embedding.
    
    Args:
        documents: List of Document objects
        chunk_size: Size of each chunk in characters
        chunk_overlap: Overlap between chunks
        
    Returns:
        List of chunked Document objects
    """
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", " ", ""]
    )
    
    chunked_docs = []
    for doc in documents:
        # Split the document
        chunks = text_splitter.split_text(doc.page_content)
        
        # Create new documents for each chunk, preserving metadata
        for i, chunk in enumerate(chunks):
            new_doc = Document(
                page_content=chunk,
                metadata={
                    **doc.metadata,
                    "chunk": i,
                    "total_chunks": len(chunks)
                }
            )
            chunked_docs.append(new_doc)
    
    return chunked_docs


def process_document_file(file_path: str) -> Tuple[List[Document], str, str]:
    """
    Complete pipeline: load document (PDF/DOCX/PPTX) and chunk it.
    
    Args:
        file_path: Path to document file
        
    Returns:
        Tuple of (chunked documents, file name, file type)
    """
    file_name = os.path.basename(file_path)
    file_ext = os.path.splitext(file_path)[1].lower()
    
    # Handle cases with multiple extensions (e.g., .pptx.pdf)
    # Check the full name for common formats
    full_name_lower = file_name.lower()
    
    # Try to load based on detected format, with fallback to PDF
    try:
        if '.pptx' in full_name_lower:
            try:
                documents, _ = load_pptx(file_path)
                file_type = "POWERPOINT"
            except Exception:
                # If PPTX fails, try PDF
                documents, _ = load_pdf(file_path)
                file_type = "PDF"
        elif '.docx' in full_name_lower:
            documents, _ = load_docx(file_path)
            file_type = "WORD"
        elif '.pdf' in full_name_lower or file_ext == '.pdf':
            documents, _ = load_pdf(file_path)
            file_type = "PDF"
        else:
            raise ValueError(f"Unsupported file type: {file_ext}. File must be PDF, DOCX, or PPTX.")
    except ValueError:
        raise
    except Exception as e:
        # Last resort: try PDF
        try:
            documents, _ = load_pdf(file_path)
            file_type = "PDF"
        except Exception as pdf_error:
            raise Exception(f"Error processing {file_name}: Could not read as any supported format. {str(e)}")
    
    chunked_docs = chunk_documents(documents)
    return chunked_docs, file_name, file_type
